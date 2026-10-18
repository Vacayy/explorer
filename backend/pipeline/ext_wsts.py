"""WSTS 월간 세계 반도체 출하액 수집 (docs/specs/value-chain.md §7, LLM 0).

출처: WSTS "Historical Billings Report" 무료 XLSX. 파일 ID가 매달 바뀌므로 안내 페이지에서 최신 .xlsx 링크를 긁는다.
'Monthly Data' 시트(실제 월간 출하액)를 쓴다 — '3MMA' 시트(3개월 이동평균)는 저장하지 않는다.
단위는 시트 머리말 "All numbers are in 1000 US$"에서 확인한다(천 달러). 머리말이 바뀌면 조용히 넘기지 않고 오류.
주의: WSTS 월은 4-4-5주 회계월이라 3·6·9·12월이 5주다 — 전월 대비는 왜곡, 전년 동월/3개월 합 비교만 의미 있다.
"""
from __future__ import annotations

import io
import re

import openpyxl
import requests

from pipeline.ext_series import upsert

PAGE_URL = "https://www.wsts.org/67/Historical-Billings-Report"
UA = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"}
UNIT = "USD thousand"
SOURCE = "WSTS Historical Billings Report"
SERIES = {"Worldwide": "wsts:world", "Americas": "wsts:americas", "Europe": "wsts:europe",
          "Japan": "wsts:japan", "Asia Pacific": "wsts:apac"}
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
          "September", "October", "November", "December"]


def parse_workbook(path_or_bytes) -> dict[str, list[tuple[str, float]]]:
    """XLSX → {series: [(YYYY-MM, 천 달러)]}. 레이아웃: 연도 행(A열 정수) 아래 지역 행, B~M열 = 1~12월"""
    src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, (bytes, bytearray)) else path_or_bytes
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    sheet = next((ws for ws in wb.worksheets if ws.title.strip().lower() == "monthly data"), None)
    if sheet is None:
        raise ValueError(f"'Monthly Data' 시트 없음: {wb.sheetnames}")
    rows = list(sheet.iter_rows(values_only=True))
    if not any(isinstance(r[0], str) and "1000 US$" in r[0] for r in rows[:10] if r):
        raise ValueError("단위 머리말('All numbers are in 1000 US$')을 찾지 못함 — 레이아웃 변경 확인 필요")
    header = next((r for r in rows[:10] if r and len(r) > 12 and list(r[1:13]) == MONTHS), None)
    if header is None:
        raise ValueError("월 머리행(January..December)을 찾지 못함 — 레이아웃 변경 확인 필요")

    out: dict[str, list[tuple[str, float]]] = {s: [] for s in SERIES.values()}
    year = None
    for r in rows:
        if not r:
            continue
        head = r[0]
        if isinstance(head, (int, float)) and 1980 <= int(head) <= 2100:
            year = int(head)
            continue
        key = SERIES.get(str(head).strip()) if head is not None else None
        if key is None or year is None:
            continue
        for m, v in enumerate(r[1:13], start=1):
            if isinstance(v, (int, float)):
                out[key].append((f"{year}-{m:02d}", float(v)))
    return out


def _latest_xlsx_url() -> str:
    res = requests.get(PAGE_URL, headers=UA, timeout=30)
    res.raise_for_status()
    links = re.findall(r'href="([^"]+\.xlsx)"', res.text, flags=re.I)
    links = [u for u in links if "billings" in u.lower()] or links
    if not links:
        raise ValueError("안내 페이지에서 .xlsx 링크를 찾지 못함")
    return requests.compat.urljoin(PAGE_URL, links[0])


def collect() -> dict:
    from database import get_connection
    try:
        url = _latest_xlsx_url()
        res = requests.get(url, headers=UA, timeout=60, allow_redirects=True)
        res.raise_for_status()
        parsed = parse_workbook(res.content)
    except Exception as e:  # noqa: BLE001 — 수집 실패는 status로 드러낸다
        return {"stored": 0, "series": [], "latest": {}, "status": "error", "error": str(e)}

    stored, series, latest = 0, [], {}
    conn = get_connection()
    try:
        for key, rows in parsed.items():
            if not rows:
                continue
            stored += upsert(conn, key, rows, UNIT, f"{SOURCE} ({url.rsplit('/', 1)[-1]})")
            series.append(key)
            latest[key] = max(rows)[0]
        conn.commit()
    finally:
        conn.close()
    missing = [k for k in SERIES.values() if k not in series]
    status = "error" if not series else ("partial" if missing else "ok")
    return {"stored": stored, "series": series, "latest": latest, "status": status,
            "unit": UNIT, "url": url, **({"missing": missing} if missing else {})}
